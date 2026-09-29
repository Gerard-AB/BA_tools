"""
Tests de la lógica de conversión.

Ejecutar con:  python -m pytest
"""

import base64
import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import app  # noqa: E402
from procesador import CABECERA, ErrorProceso, clasificar, convertir  # noqa: E402

EJEMPLO = Path(__file__).resolve().parents[1] / "ejemplo" / "inscripciones-ejemplo.csv"

CSV_MINIMO = (
    "Beneficiari;Classe;Total compra\n"
    "Gómez Puig, Ada;Quart de Primària-A;45\n"
)


def csv_bytes(texto):
    return texto.encode("utf-8")


# --- Clasificación de cursos -------------------------------------------------

@pytest.mark.parametrize("curso, hoja", [
    ("Tercer de Primària-A", "Grumets"),
    ("Quart de Primària-C", "Grumets"),
    ("Cinqué de Primària-B", "Ulls Oberts"),
    ("Sisè de Primària", "Ulls Oberts"),
    ("Primer de la ESO-A", "Mà Oberta"),
    ("Segón de la ESO", "Mà Oberta"),
    ("Segon d'ESO-B", "Mà Oberta"),
    ("Tercer de la ESO-A", "Cor Obert I - II"),
    ("Quart de la ESO", "Cor Obert I - II"),
    ("Quart d'ESO-B", "Cor Obert I - II"),
    # El ordinal solo no basta: Tercer de Primària y Tercer de la ESO
    # comparten nivel pero van a grupos distintos.
    ("Primer de Primària", "Cor Obert III - IV"),
    ("Segón de Primària", "Cor Obert III - IV"),
    ("Infantil 5 anys", "Cor Obert III - IV"),
    ("", "Cor Obert III - IV"),
])
def test_clasificar(curso, hoja):
    assert clasificar(curso)["hoja"] == hoja


# --- Conversión --------------------------------------------------------------

def test_separa_nombre_y_apellidos():
    hojas, _ = convertir(csv_bytes(CSV_MINIMO), "x.csv")
    assert hojas[0]["filas"][0][:3] == ["Ada", "Gómez Puig", "Quart de Primària-A"]


def test_sin_coma_todo_va_al_nombre_y_avisa():
    hojas, avisos = convertir(
        csv_bytes("Beneficiari;Classe;Total compra\nSinComa;Quart de Primària;45\n"), "x.csv")
    assert hojas[0]["filas"][0][:2] == ["SinComa", ""]
    assert len(avisos) == 1


@pytest.mark.parametrize("bruto, esperado", [
    ("45", 45), ("54,50", 54.5), ("1.234,56 €", 1234.56), ("66,5", 66.5),
])
def test_importes(bruto, esperado):
    hojas, _ = convertir(
        csv_bytes(f"Beneficiari;Classe;Total compra\nApe, Nom;Quart de Primària;{bruto}\n"), "x.csv")
    assert hojas[0]["filas"][0][6] == esperado


def test_importe_no_numerico_se_copia_y_avisa():
    hojas, avisos = convertir(
        csv_bytes("Beneficiari;Classe;Total compra\nApe, Nom;Quart de Primària;gratis\n"), "x.csv")
    assert hojas[0]["filas"][0][6] == "gratis"
    assert len(avisos) == 1


def test_cabeceras_con_acentos_mayusculas_y_alias():
    hojas, _ = convertir(
        csv_bytes("BENEFICIARIO;Curso;Importe\nApe, Nom;Quart de Primària;45\n"), "x.csv")
    assert hojas[0]["filas"][0][0] == "Nom"


def test_csv_y_xlsx_dan_el_mismo_resultado():
    crudo = EJEMPLO.read_bytes()
    desde_csv, _ = convertir(crudo, "ejemplo.csv")

    wb = Workbook()
    ws = wb.active
    for linea in crudo.decode("utf-8-sig").splitlines():
        ws.append(linea.split(";"))
    buffer = io.BytesIO()
    wb.save(buffer)
    desde_xlsx, _ = convertir(buffer.getvalue(), "ejemplo.xlsx")

    assert [h["filas"] for h in desde_csv] == [h["filas"] for h in desde_xlsx]


# --- Errores esperados -------------------------------------------------------

@pytest.mark.parametrize("contenido, nombre, trozo", [
    (b"", "x.csv", "vac"),
    (b"a;b;c\n1;2;3\n", "x.csv", "No se ha encontrado la columna"),
    (b"lo que sea", "x.pdf", "Formato no soportado"),
    (b"lo que sea", "x.xls", "Excel antiguo"),
])
def test_errores(contenido, nombre, trozo):
    with pytest.raises(ErrorProceso, match=trozo):
        convertir(contenido, nombre)


# --- XLSX generado -----------------------------------------------------------

def test_un_xlsx_por_grupo_con_su_color():
    hojas, _ = convertir(EJEMPLO.read_bytes(), "ejemplo.csv")

    assert [h["nombre"] for h in hojas] == [
        "Grumets", "Ulls Oberts", "Mà Oberta", "Cor Obert I - II", "Cor Obert III - IV"]

    for hoja in hojas:
        wb = load_workbook(io.BytesIO(hoja["xlsx"]))
        # Cada fichero lleva una sola hoja: la de su grupo.
        assert wb.sheetnames == [hoja["nombre"]]

        ws = wb.active
        assert [c.value for c in ws[1]] == CABECERA
        assert ws.max_row == len(hoja["filas"]) + 1
        assert ws.cell(1, 1).fill.fgColor.rgb[-6:] == hoja["color"]
        assert ws.freeze_panes == "A2"
        assert "€" in ws.cell(2, len(CABECERA)).number_format


def test_nombre_de_archivo_por_grupo():
    hojas, _ = convertir(EJEMPLO.read_bytes(), "inscripciones-ejemplo.csv")

    assert [h["archivo"] for h in hojas] == [
        "inscripciones-ejemplo_grumets.xlsx",
        "inscripciones-ejemplo_ulls-oberts.xlsx",
        "inscripciones-ejemplo_ma-oberta.xlsx",
        "inscripciones-ejemplo_cor-obert-i-ii.xlsx",
        "inscripciones-ejemplo_cor-obert-iii-iv.xlsx",
    ]


# --- API ---------------------------------------------------------------------

def test_endpoint_devuelve_un_excel_por_grupo():
    cliente = TestClient(app)
    respuesta = cliente.post(
        "/procesar", files={"fichero": ("ejemplo.csv", EJEMPLO.read_bytes(), "text/csv")})

    assert respuesta.status_code == 200
    datos = respuesta.json()
    assert datos["total_filas"] == 22
    assert len(datos["hojas"]) == 5

    for hoja in datos["hojas"]:
        assert hoja["archivo"].startswith("ejemplo_") and hoja["archivo"].endswith(".xlsx")
        wb = load_workbook(io.BytesIO(base64.b64decode(hoja["xlsx"])))
        assert wb.sheetnames == [hoja["nombre"]]
        assert wb.active.max_row == len(hoja["filas"]) + 1


def test_endpoint_rechaza_fichero_invalido():
    cliente = TestClient(app)
    respuesta = cliente.post(
        "/procesar", files={"fichero": ("x.pdf", b"algo", "application/pdf")})

    assert respuesta.status_code == 400
    assert "error" in respuesta.json()


def test_pagina_y_logo_se_sirven():
    cliente = TestClient(app)
    assert cliente.get("/").status_code == 200
    assert cliente.get("/logo.png").headers["content-type"] == "image/png"
